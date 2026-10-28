"""Second-annotator kit for the 59 FHIR pairs (41 dev + 18 confirm), blind to the draft GT, plus agreement scoring.

  python annotation_kit.py make             -> annotation/anotasi_kedua.xlsx (one row per source column)
  python annotation_kit.py score FILE.xlsx  -> agreement with the draft GT + disagreement list for adjudication

Agreement is reported at two levels:
  resource level  sheet "Resource": does the annotator accept the FHIR resource assigned to each table (raw agreement)
  element level   per source column, label = target element (1:1 or n:1) or NONE; Cohen's kappa over those labels
                  ("ambigu" rows reported apart), plus P/R/F1 of annotator 2's 1:1 links vs. the draft gt.
"""
import json
import pathlib
import sys
from collections import Counter

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from hybrid import SCHEMAS, columns

PAIRS = [pathlib.Path("data/subset_e/pairs_e.json"), pathlib.Path("data/subset_e/pairs_confirm.json")]
OUT = pathlib.Path("annotation/anotasi_kedua.xlsx")
RELATIONS = ["1:1", "n:1", "ambigu", "tidak ada"]
GUIDE = """Tugas: untuk setiap kolom tabel sumber (SIMRS/aplikasi kesehatan Indonesia), pilih elemen FHIR R4 yang paling tepat di resource target yang sudah ditentukan, atau "tidak ada".

Kolom yang diisi: "Elemen target" (dropdown) dan "Relasi".
- 1:1: kolom ini adalah padanan utama elemen tersebut. Satu elemen hanya boleh punya SATU kolom 1:1 per tabel.
- n:1: kolom ini juga jatuh ke elemen yang sudah dipakai kolom 1:1 lain (mis. jam_reg dan tgl_registrasi -> period; pilih yang paling representatif sebagai 1:1).
- ambigu: Anda ragu atau dua tafsiran sama kuat. Tulis alasannya di Catatan.
- tidak ada: tidak ada padanan di resource ini (biarkan Elemen target kosong).

Konvensi:
1. FHIR R4 DASAR saja, tanpa ekstensi SATUSEHAT. Tempat lahir, agama, nama ibu kandung, kewarganegaraan pada Patient = "tidak ada".
2. Elemen tidak dipecah ke komponen datatype: pakai "name", bukan "name.family"; "address", bukan "address.city".
3. Kode vs nama untuk atribut yang sama (kd_kec vs nm_kec) -> pilih satu sebagai 1:1, sisanya n:1.
4. FK yang hanya terhubung lewat join ke tabel lain (bukan referensi langsung ke resource) -> ambigu.
5. Nomor kunjungan/rawat (no_rawat, no_redis, idRekamMedis) pada tabel selain kunjungan itu sendiri adalah referensi -> "encounter", bukan "identifier".
6. Tabel tanda vital: kolom tekanan darah -> component.value[x] (1:1), vital lain -> component.value[x] (n:1).
7. Kolom tarif/harga dan kolom audit aplikasi (created_at, user input) -> tidak ada, kecuali jelas padanannya.

Di sheet "Resource", nyatakan apakah Anda setuju dengan resource FHIR yang ditetapkan untuk tiap tabel; bila tidak, tulis resource alternatif. Kolom pada sheet "Anotasi" tetap diisi terhadap resource yang ditetapkan.

Jangan melihat file GT lain dan jangan menjalankan model. Perkiraan waktu ±4 jam; boleh dicicil.
"""


def load_pairs() -> list[dict]:
    return [p for f in PAIRS for p in json.loads(f.read_text()) if p["target"].startswith("fhir__")]


def make():
    wb = Workbook()
    ws = wb.active
    ws.title = "Petunjuk"
    ws["A1"] = GUIDE
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140
    ws.row_dimensions[1].height = 420
    pairs = load_pairs()
    ann = wb.create_sheet("Anotasi")
    head = ["pasangan", "tabel sumber", "kolom sumber", "tipe", "resource target", "Elemen target", "Relasi", "Catatan"]
    ann.append(head)
    for c in ann[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (10, 34, 26, 16, 18, 30, 12, 40)):
        ann.column_dimensions[col].width = w
    ann.freeze_panes = "A2"
    rel = DataValidation(type="list", formula1=f'"{",".join(RELATIONS)}"', allow_blank=True)
    ann.add_data_validation(rel)
    fill = PatternFill("solid", fgColor="FFF2CC")
    ref_sheets = {}
    for p in pairs:
        res = p["target"].removeprefix("fhir__")
        if res not in ref_sheets:
            rs = wb.create_sheet(f"FHIR_{res}"[:31])
            rs.append(["elemen", "tipe", "deskripsi (FHIR short)"])
            for c in json.loads((SCHEMAS / f"{p['target']}.json").read_text())["columns"]:
                rs.append([c["name"], c["type"], c["desc"]])
            rs.column_dimensions["A"].width, rs.column_dimensions["C"].width = 32, 90
            ref_sheets[res] = (rs.title, rs.max_row)
            dv = DataValidation(type="list", formula1=f"='{rs.title}'!$A$2:$A${rs.max_row}", allow_blank=True)
            ann.add_data_validation(dv)
            ref_sheets[res] += (dv,)
        schema = json.loads((SCHEMAS / f"{p['source']}.json").read_text())["columns"]
        for c in schema:
            ann.append([p["id"], p["source"], c["name"], c["type"], res, None, None, None])
            r = ann.max_row
            ref_sheets[res][2].add(f"F{r}")
            rel.add(f"G{r}")
            ann[f"F{r}"].fill = ann[f"G{r}"].fill = fill
    res = wb.create_sheet("Resource", 1)
    res.append(["pasangan", "tabel sumber", "resource yang ditetapkan", "Setuju? (ya/tidak)", "Resource alternatif", "Catatan"])
    for c in res[1]:
        c.font = Font(bold=True)
    yn = DataValidation(type="list", formula1='"ya,tidak"', allow_blank=True)
    res.add_data_validation(yn)
    for p in pairs:
        res.append([p["id"], p["source"], p["target"].removeprefix("fhir__"), None, None, None])
        yn.add(f"D{res.max_row}")
        res[f"D{res.max_row}"].fill = fill
    for col, w in zip("ABCDEF", (10, 34, 22, 18, 22, 40)):
        res.column_dimensions[col].width = w
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"{OUT}: {len(pairs)} pairs, {ann.max_row - 1} source columns")


def kappa(a: list, b: list) -> float:
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in ca) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def score(path: str):
    df = pd.read_excel(path, sheet_name="Anotasi").rename(columns={"Elemen target": "el", "Relasi": "rel"})
    assert df["rel"].notna().all(), f"{df['rel'].isna().sum()} rows without Relasi"
    gt = {p["id"]: p for p in load_pairs()}
    rows = []
    for _, r in df.iterrows():
        p = gt[r["pasangan"]]
        g1 = {s: t for s, t in p["gt"] + p["gt_n1"]}
        amb = {s for s, *_ in p["ambiguous"]}
        draft = "AMBIGU" if r["kolom sumber"] in amb else g1.get(r["kolom sumber"], "NONE")
        second = "AMBIGU" if r["rel"] == "ambigu" else ("NONE" if r["rel"] == "tidak ada" else str(r["el"]))
        rows.append({"pair": r["pasangan"], "col": r["kolom sumber"], "draft": draft, "second": second,
                     "rel2": r["rel"], "note": r.get("Catatan")})
    d = pd.DataFrame(rows)
    clear = d[(d.draft != "AMBIGU") & (d.second != "AMBIGU")]
    print(f"columns {len(d)}; both non-ambiguous {len(clear)}; raw agreement {(clear.draft == clear.second).mean():.3f}; "
          f"Cohen's kappa {kappa(clear.draft.tolist(), clear.second.tolist()):.3f}")
    a2 = {(r.pair, r.col, r.second) for r in d.itertuples() if r.rel2 == "1:1"}
    a1 = {(pid, s, t) for pid, p in gt.items() for s, t in p["gt"]}
    tp = len(a1 & a2)
    print(f"1:1 links: draft {len(a1)}, second {len(a2)}, shared {tp}; "
          f"P {tp / max(len(a2), 1):.3f} R {tp / max(len(a1), 1):.3f} F1 {2 * tp / max(len(a1) + len(a2), 1):.3f}")
    r = pd.read_excel(path, sheet_name="Resource")
    agreed = r["Setuju? (ya/tidak)"].astype(str).str.strip().str.lower().eq("ya")
    print(f"resource level: annotator 2 accepts the assigned resource for {agreed.sum()}/{len(r)} tables; "
          f"alternatives: {r.loc[~agreed, 'Resource alternatif'].dropna().tolist()}")
    out = pathlib.Path(path).with_name("ketidaksepakatan.csv")
    d[d.draft != d.second].to_csv(out, index=False)
    print(f"disagreements for adjudication -> {out}")


def demo():
    assert kappa(list("aabb"), list("aabb")) == 1.0
    assert abs(kappa(list("ab"), list("ba")) + 1.0) < 1e-9


if __name__ == "__main__":
    demo()
    {"make": lambda: make(), "score": lambda: score(sys.argv[2])}[sys.argv[1]]()
